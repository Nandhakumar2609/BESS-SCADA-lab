#!/usr/bin/env python3
"""Generate a Power BI-ready data set from the lab's physics model and alarm engine.

Runs the same SiteModel the Modbus simulator uses, decodes every value from the device
register images (exactly what the gateway would read), applies a daily arbitrage dispatch
(charge on the midday solar peak, discharge on the evening peak) and runs the real alarm
engine. Output: powerbi/BESS_PowerBI_Data.xlsx with Excel tables Power BI imports directly.

    python tools/export_powerbi.py --days 7
"""
import argparse
import sys
from datetime import datetime, timedelta
from pathlib import Path

import pandas as pd

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from bess_lab.config import CONFIG_DIR, load_site, load_yaml  # noqa: E402
from bess_lab.gateway.alarms import AlarmEngine  # noqa: E402
from bess_lab.gateway.poller import DeviceResult  # noqa: E402
from bess_lab.regmap import RegisterMap  # noqa: E402
from bess_lab.sim.model import SiteModel  # noqa: E402

STATE = {0: "STOPPED", 1: "STANDBY", 2: "CHARGING", 3: "DISCHARGING", 4: "FAULT"}


def schedule_kw(t: datetime, soc: float) -> float:
    """Simple energy-arbitrage schedule (kW, + discharge)."""
    h = t.hour + t.minute / 60
    if 10 <= h < 15 and soc < 92:
        return -450.0             # absorb midday solar, cheap prices
    if 18 <= h < 21.5 and soc > 12:
        return 600.0              # evening peak
    if 6.5 <= h < 8.5 and soc > 45:
        return 350.0              # morning peak
    return 0.0


def main(days: int, seed: int, out: Path) -> None:
    site = load_site()
    regmap = RegisterMap.load(CONFIG_DIR / "register_map.csv")
    devices = site["devices"]
    model = SiteModel(regmap, devices, site["site"], speed=1.0, seed=seed)
    model.fault_rate_per_h = 0.03
    start = datetime(2026, 10, 5, 0, 0)          # Monday
    model.sim_time = start.timestamp()
    model.soc = 0.40
    engine = AlarmEngine(load_yaml("alarms.yaml"), devices, regmap, stale_after_s=600)
    pcs = model.images["pcs01"]

    rows, events, fault_since = [], [], None
    steps = days * 24 * 60
    for i in range(steps):
        now = start + timedelta(minutes=i)
        # SCADA dispatch + supervised auto-reset (after >= 1 min in FAULT)
        soc_now = model.soc * 100
        pcs.set("p_setpoint_kw", schedule_kw(now, soc_now))
        model.step(60.0)
        vals = {d.name: {p.name: p.decode((model.images[d.name].input if p.table == "input" else
                                            model.images[d.name].holding)[p.address:p.end])
                         for p in regmap.select(d.profile)} for d in devices}
        results = {n: DeviceResult(n, d.profile, ok=True, values=vals[n]) for n, d in zip(vals, devices)}
        for e in engine.evaluate(results, now.timestamp()):
            events.append(e)
        p, b, h, poi, aux = (vals[k] for k in ("pcs01", "bms01", "hvac01", "meter_poi", "meter_aux"))
        state = int(p["state"])
        if state == 4:
            fault_since = fault_since or now
            if now - fault_since >= timedelta(minutes=1):
                pcs.set("fault_reset", 1)
        else:
            fault_since = None
        available = int(state != 4 and int(b["status"]) != 4 and int(b["contactor_closed"]) == 1)
        rows.append({
            "Timestamp": now, "Date": now.date(), "Hour": now.hour,
            "SoC_pct": round(b["soc_pct"], 1), "SoH_pct": round(b["soh_pct"], 1),
            "PCS_State": STATE[state], "PCS_Power_kW": round(p["active_power_kw"], 1),
            "PCS_Reactive_kvar": round(p["reactive_power_kvar"], 1), "POI_Power_kW": round(poi["p_total_kw"], 1),
            "AUX_Power_kW": round(aux["p_total_kw"], 2), "HVAC_Power_kW": round(h["power_kw"], 2),
            "Grid_Frequency_Hz": round(p["frequency_hz"], 3), "AC_Voltage_V": round(p["ac_voltage_v"], 1),
            "Pack_Voltage_V": round(b["pack_voltage_v"], 1), "Pack_Current_A": round(b["pack_current_a"], 1),
            "Cell_Vmax_V": b["max_cell_voltage_v"], "Cell_Vmin_V": b["min_cell_voltage_v"],
            "Cell_Tmax_C": b["max_cell_temp_c"], "Cell_Tmin_C": b["min_cell_temp_c"],
            "Container_Temp_C": h["container_temp_c"], "Ambient_Temp_C": h["ambient_temp_c"],
            "IGBT_Temp_C": p["igbt_temp_c"], "Humidity_pct": h["humidity_pct"],
            "Compressor_On": int(h["compressor_on"]), "Available": available,
            # energy in this minute (kWh) - additive, so Power BI can SUM it
            "Energy_Discharged_kWh": round(max(p["active_power_kw"], 0) / 60, 4),
            "Energy_Charged_kWh": round(max(-p["active_power_kw"], 0) / 60, 4),
            "POI_Export_kWh": round(max(poi["p_total_kw"], 0) / 60, 4),
            "POI_Import_kWh": round(max(-poi["p_total_kw"], 0) / 60, 4),
            "AUX_Energy_kWh": round(aux["p_total_kw"] / 60, 4),
        })
        if i % 1440 == 0:
            print(f"day {i // 1440 + 1}/{days}  SoC {b['soc_pct']:.1f}%")

    tel = pd.DataFrame(rows)

    # alarm events -> one row per occurrence (raised/cleared paired)
    open_, alarms = {}, []
    for e in events:
        key = (e.device, e.alarm_id)
        if e.state == "RAISED":
            open_[key] = e
        elif key in open_:
            r = open_.pop(key)
            alarms.append((r, e))
    for key, r in open_.items():
        alarms.append((r, None))
    alarm_df = pd.DataFrame([{
        "Alarm_ID": r.alarm_id, "Device": r.device, "Severity": r.severity, "Message": r.message,
        "Raised": datetime.fromtimestamp(r.ts), "Cleared": datetime.fromtimestamp(c.ts) if c else None,
        "Duration_min": round((c.ts - r.ts) / 60, 1) if c else None,
        "Date": datetime.fromtimestamp(r.ts).date(),
    } for r, c in sorted(alarms, key=lambda x: x[0].ts)])

    daily = tel.groupby("Date").agg(
        Availability_pct=("Available", lambda s: round(100 * s.mean(), 2)),
        Energy_Charged_kWh=("Energy_Charged_kWh", "sum"), Energy_Discharged_kWh=("Energy_Discharged_kWh", "sum"),
        POI_Export_kWh=("POI_Export_kWh", "sum"), POI_Import_kWh=("POI_Import_kWh", "sum"),
        AUX_Energy_kWh=("AUX_Energy_kWh", "sum"), SoC_min_pct=("SoC_pct", "min"), SoC_max_pct=("SoC_pct", "max"),
        Cell_Tmax_C=("Cell_Tmax_C", "max"), Fault_Minutes=("PCS_State", lambda s: int((s == "FAULT").sum())),
    ).reset_index()
    for c in ("Energy_Charged_kWh", "Energy_Discharged_kWh", "POI_Export_kWh", "POI_Import_kWh", "AUX_Energy_kWh"):
        daily[c] = daily[c].round(1)
    daily["Equivalent_Cycles"] = (daily["Energy_Discharged_kWh"] / site["site"]["capacity_kwh"]).round(3)
    daily["Alarm_Count"] = daily["Date"].map(alarm_df.groupby("Date").size() if len(alarm_df) else {}).fillna(0).astype(int)

    dev_df = pd.DataFrame([{"Device": d.name, "Type": d.profile.upper(), "IP_Address": d.host, "Port": d.port,
                            "Unit_ID": d.unit_id, "Zone": "Field (VLAN 10)"} for d in devices])
    map_df = pd.DataFrame([{"Profile": p.profile, "Table": p.table, "Reference": p.reference, "Address": p.address,
                            "Point": p.name, "Data_Type": p.data_type, "Byte_Order": p.byte_order, "Scale": p.scale,
                            "Unit": p.unit, "Access": p.access, "Description": p.description} for p in regmap.points])
    dates = pd.DataFrame({"Date": pd.date_range(start.date(), periods=days).date})
    dates["Day"] = [d.strftime("%a %d %b") for d in dates["Date"]]
    dates["Weekday"] = [d.strftime("%A") for d in dates["Date"]]

    write_xlsx(out, {"Telemetry": tel, "Alarms": alarm_df, "Daily_KPI": daily, "Devices": dev_df,
                     "Dates": dates, "Register_Map": map_df}, site, days)
    print(f"wrote {out}: {len(tel)} telemetry rows, {len(alarm_df)} alarms, availability "
          f"{100 * tel['Available'].mean():.2f}%")


def write_xlsx(path, sheets, site, days):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    wb = Workbook()
    ws = wb.active
    ws.title = "README"
    head = Font(name="Arial", bold=True, size=14)
    body = Font(name="Arial", size=10)
    lines = [
        ("BESS Site 01 - SCADA data set for Power BI", head),
        (f"{site['site']['name']}  |  {days} days at 1-minute resolution", body),
        ("Generated by tools/export_powerbi.py from the lab's simulator model, register map and alarm engine.", body),
        ("Values are decoded from the simulated Modbus register images, i.e. what the SCADA gateway reads.", body),
        ("", body),
        ("Tables (each sheet holds one Excel table that Power BI lists by name):", Font(name="Arial", bold=True)),
        ("tblTelemetry   1-minute site data: SoC, power, temperatures, state, availability, energy per minute", body),
        ("tblAlarms      one row per alarm occurrence: raised, cleared, duration, severity", body),
        ("tblDailyKPI    pre-aggregated daily KPIs (availability, throughput, cycles, alarms)", body),
        ("tblDevices     device list with IP, port, unit ID", body),
        ("tblDates       date dimension - relate tblDates[Date] to the Date column of the other tables", body),
        ("tblRegisterMap Modbus protocol mapping table", body),
        ("", body),
        ("Sign convention: PCS and POI power are + discharge / export, - charge / import (kW).", body),
        ("Energy columns are kWh per minute row, so SUM gives energy over any period.", body),
        ("Build steps and DAX measures: powerbi/POWERBI_GUIDE.md", body),
    ]
    for i, (t, f) in enumerate(lines, start=1):
        ws.cell(row=i, column=1, value=t).font = f
    ws.column_dimensions["A"].width = 110

    hdr_fill = PatternFill("solid", fgColor="2F5D86")
    for name, df in sheets.items():
        ws = wb.create_sheet(name)
        cols = list(df.columns)
        ws.append(cols)
        for row in df.itertuples(index=False):
            ws.append([None if (isinstance(v, float) and v != v) else v for v in row])
        for j, c in enumerate(cols, start=1):
            cell = ws.cell(row=1, column=j)
            cell.font = Font(name="Arial", bold=True, color="FFFFFF")
            cell.fill = hdr_fill
            cell.alignment = Alignment(horizontal="center")
            ws.column_dimensions[get_column_letter(j)].width = max(11, min(42, len(c) + 3))
            fmt = None
            if c in ("Timestamp", "Raised", "Cleared"):
                fmt, ws.column_dimensions[get_column_letter(j)].width = "yyyy-mm-dd hh:mm", 18
            elif c == "Date":
                fmt = "yyyy-mm-dd"
            if fmt:
                for r in range(2, len(df) + 2):
                    ws.cell(row=r, column=j).number_format = fmt
        ref = f"A1:{get_column_letter(len(cols))}{len(df) + 1}"
        tbl = Table(displayName="tbl" + name.replace("_", ""), ref=ref)
        tbl.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
        ws.add_table(tbl)
        ws.freeze_panes = "A2"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--days", type=int, default=7)
    ap.add_argument("--seed", type=int, default=11)
    ap.add_argument("--out", default=str(ROOT / "powerbi" / "BESS_PowerBI_Data.xlsx"))
    a = ap.parse_args()
    main(a.days, a.seed, Path(a.out))
